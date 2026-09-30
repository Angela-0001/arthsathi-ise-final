"""
OCR module — extracts text from images, PDFs, and Word documents.
- Images: Tesseract OCR
- PDFs: PyMuPDF (text extraction + image OCR fallback)
- Word (.docx): python-docx
"""
from PIL import Image
import io
import sys


def extract_text(image_bytes: bytes, lang: str = "hin+eng") -> str:
    """Auto-detect file type and extract text accordingly."""
    # Detect file type by magic bytes
    if image_bytes[:4] == b'%PDF':
        return _extract_pdf(image_bytes, lang)
    elif image_bytes[:4] == b'PK\x03\x04':
        return _extract_docx(image_bytes)
    else:
        return _extract_image(image_bytes, lang)


def _extract_image(image_bytes: bytes, lang: str) -> str:
    try:
        import pytesseract

        if sys.platform == "win32":
            import os
            path = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
            if os.path.exists(path):
                pytesseract.pytesseract.tesseract_cmd = path

        image = Image.open(io.BytesIO(image_bytes)).convert("L")

        try:
            text = pytesseract.image_to_string(image, lang=lang)
        except Exception:
            text = pytesseract.image_to_string(image, lang="eng")

        return text.strip()

    except Exception as e:
        print(f"[OCR] Image error: {e}")
        return ""


def _extract_pdf(pdf_bytes: bytes, lang: str) -> str:
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        full_text = []

        for page in doc:
            # Try direct text extraction first (fast, works for digital PDFs)
            text = page.get_text("text").strip()
            if text and len(text) > 20:
                full_text.append(text)
            else:
                # Scanned PDF — render page as image and OCR it
                mat = fitz.Matrix(2.0, 2.0)  # 2x zoom for better quality
                pix = page.get_pixmap(matrix=mat)
                img_bytes = pix.tobytes("png")
                ocr_text = _extract_image(img_bytes, lang)
                if ocr_text:
                    full_text.append(ocr_text)

        doc.close()
        return "\n".join(full_text).strip()

    except ImportError:
        print("[OCR] PyMuPDF not installed. Run: pip install pymupdf")
        return ""
    except Exception as e:
        print(f"[OCR] PDF error: {e}")
        return ""


def _extract_docx(docx_bytes: bytes) -> str:
    try:
        from docx import Document

        doc = Document(io.BytesIO(docx_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        return "\n".join(paragraphs).strip()

    except ImportError:
        print("[OCR] python-docx not installed. Run: pip install python-docx")
        return ""
    except Exception as e:
        print(f"[OCR] Word error: {e}")
        return ""
